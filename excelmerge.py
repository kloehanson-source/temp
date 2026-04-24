#!/usr/bin/env python3
"""
excelmerge.py - Consolidate multiple Excel files into a single output file.
Usage: python excelmerge.py "path/to/folder"
"""

import sys
import os
import subprocess
from datetime import datetime


# ---------------------------------------------------------------------------
# Auto-install dependencies
# ---------------------------------------------------------------------------

def _ensure_packages():
    needed = [("openpyxl", "openpyxl"), ("xlrd", "xlrd")]
    for pkg, mod in needed:
        try:
            __import__(mod)
        except ImportError:
            print(f"Package '{pkg}' not found. Installing...")
            subprocess.check_call(
                [sys.executable, "-m", "pip", "install", pkg, "--quiet"],
                stderr=subprocess.DEVNULL,
            )
            print(f"  '{pkg}' installed successfully.")

_ensure_packages()

import openpyxl          # noqa: E402  (imported after install check)
import xlrd              # noqa: E402


# ---------------------------------------------------------------------------
# File discovery
# ---------------------------------------------------------------------------

def _find_excel_files(folder: str) -> list:
    results = []
    for fname in sorted(os.listdir(folder)):
        fpath = os.path.join(folder, fname)
        if not os.path.isfile(fpath):
            continue
        lower = fname.lower()
        # Skip any previously generated output files so re-runs stay clean
        if lower.startswith("consolidated_") and lower.endswith(".xlsx"):
            continue
        if lower.endswith(".xlsx") or lower.endswith(".xls"):
            results.append(fpath)
    return results


# ---------------------------------------------------------------------------
# Header reading (one pass, minimal memory)
# ---------------------------------------------------------------------------

def _headers_xlsx(filepath: str):
    """Return (list[str], None) on success or (None, err_str) on failure."""
    try:
        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
        ws = wb.active
        headers = []
        for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
            headers = [str(c).strip() if c is not None else "" for c in row]
        wb.close()
        return headers, None
    except Exception as exc:
        return None, str(exc)


def _headers_xls(filepath: str):
    """Return (list[str], None) on success or (None, err_str) on failure."""
    try:
        wb = xlrd.open_workbook(filepath, on_demand=True)
        ws = wb.sheet_by_index(0)
        if ws.nrows == 0:
            return [], None
        headers = [str(ws.cell(0, c).value).strip() for c in range(ws.ncols)]
        wb.release_resources()
        return headers, None
    except Exception as exc:
        return None, str(exc)


def _collect_all_columns(files: list) -> list:
    """
    Scan every file for its first-row headers.
    Return an ordered list of unique column display-names (case-insensitive dedup,
    preserving the first-seen capitalisation).
    """
    seen: dict = {}   # lowercase key -> display name
    for filepath in files:
        fname = os.path.basename(filepath)
        ext = os.path.splitext(filepath)[1].lower()
        if ext == ".xlsx":
            headers, err = _headers_xlsx(filepath)
        else:
            headers, err = _headers_xls(filepath)

        if headers is None:
            print(f"  Warning: could not read headers from '{fname}': {err}")
            continue

        for col in headers:
            if not col:
                continue
            key = col.lower()
            if key not in seen:
                seen[key] = col

    return list(seen.values())


# ---------------------------------------------------------------------------
# Row streaming
# ---------------------------------------------------------------------------

def _stream_xlsx(filepath: str, col_map: dict, n_out: int, filters: list):
    """Yield output rows from an xlsx file using read-only / streaming mode."""
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    ws = wb.active
    file_headers_lower = None
    try:
        for row_vals in ws.iter_rows(values_only=True):
            # First row is the header
            if file_headers_lower is None:
                file_headers_lower = [
                    str(c).strip().lower() if c is not None else ""
                    for c in row_vals
                ]
                continue

            # Skip rows that are completely empty
            if all(c is None for c in row_vals):
                continue

            out_row = [None] * n_out
            for fi, fh in enumerate(file_headers_lower):
                if fi < len(row_vals) and fh in col_map:
                    val = row_vals[fi]
                    out_row[col_map[fh]] = val

            if _passes_filters(out_row, filters):
                yield out_row
    finally:
        wb.close()


def _stream_xls(filepath: str, col_map: dict, n_out: int, filters: list):
    """Yield output rows from an xls file using xlrd."""
    wb = xlrd.open_workbook(filepath, on_demand=True)
    ws = wb.sheet_by_index(0)

    if ws.nrows < 2:
        return

    file_headers_lower = [
        str(ws.cell(0, c).value).strip().lower() for c in range(ws.ncols)
    ]

    for ri in range(1, ws.nrows):
        out_row = [None] * n_out
        all_empty = True

        for fi, fh in enumerate(file_headers_lower):
            if fh not in col_map:
                continue
            cell = ws.cell(ri, fi)
            if cell.ctype == xlrd.XL_CELL_EMPTY:
                val = None
            elif cell.ctype == xlrd.XL_CELL_DATE:
                try:
                    val = xlrd.xldate_as_datetime(cell.value, wb.datemode)
                except Exception:
                    val = cell.value
            else:
                val = cell.value
            out_row[col_map[fh]] = val
            if val is not None and val != "":
                all_empty = False

        if all_empty:
            continue

        if _passes_filters(out_row, filters):
            yield out_row


# ---------------------------------------------------------------------------
# Filtering
# ---------------------------------------------------------------------------

def _passes_filters(row: list, filters: list) -> bool:
    """
    filters is a list of (col_idx, [lowercase_partial_values]).
    A row passes when every filter has at least one matching value (OR within
    a filter, AND across filters).
    """
    for col_idx, values in filters:
        cell_str = str(row[col_idx]).lower() if row[col_idx] is not None else ""
        if not any(v in cell_str for v in values):
            return False
    return True


# ---------------------------------------------------------------------------
# Interactive prompts
# ---------------------------------------------------------------------------

def _prompt_exclude_columns(all_columns: list) -> list:
    print("\nUnique columns found across all files:")
    for i, col in enumerate(all_columns, 1):
        print(f"  {i:4}. {col}")

    raw = input(
        "\nEnter column numbers to EXCLUDE (comma-separated), or press Enter to keep all: "
    ).strip()

    if not raw:
        return list(all_columns)

    try:
        exclude_nums = {int(x.strip()) for x in raw.split(",") if x.strip().isdigit()}
    except ValueError:
        print("Could not parse input — keeping all columns.")
        return list(all_columns)

    included = [col for i, col in enumerate(all_columns, 1) if i not in exclude_nums]
    excluded  = [col for i, col in enumerate(all_columns, 1) if i in exclude_nums]
    if excluded:
        print(f"Excluding {len(excluded)} column(s): {', '.join(excluded)}")
    return included


def _prompt_filters(included_cols: list) -> list:
    """
    Return a list of (col_idx, [lowercase_partial_values]) tuples.
    col_idx is the index into included_cols (i.e. the output row index).
    """
    filters = []

    while True:
        choice = input("\nFilter by a column value? (y/n): ").strip().lower()
        if choice != "y":
            break

        print("\nIncluded columns:")
        for i, col in enumerate(included_cols, 1):
            print(f"  {i:4}. {col}")

        col_raw = input("Enter column number to filter on: ").strip()
        if not col_raw.isdigit():
            print("Invalid — skipping.")
            continue

        col_num = int(col_raw)
        if col_num < 1 or col_num > len(included_cols):
            print(f"Out of range (1–{len(included_cols)}) — skipping.")
            continue

        col_idx = col_num - 1
        col_name = included_cols[col_idx]

        vals_raw = input(
            f"Enter value(s) to keep in '{col_name}' (comma-separated, partial match OK): "
        ).strip()
        if not vals_raw:
            print("No values entered — skipping.")
            continue

        filter_vals = [v.strip().lower() for v in vals_raw.split(",") if v.strip()]
        filters.append((col_idx, filter_vals))
        print(f"  Filter added: '{col_name}' contains any of {filter_vals}")

        more = input("Add another filter? (y/n): ").strip().lower()
        if more != "y":
            break

    return filters


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    if len(sys.argv) < 2:
        print("Usage: python excelmerge.py <folder_path>")
        sys.exit(1)

    folder = sys.argv[1]
    if not os.path.isdir(folder):
        print(f"Error: '{folder}' is not a valid directory.")
        sys.exit(1)

    # ---- Discover files ------------------------------------------------
    print(f"\nScanning folder: {folder}")
    files = _find_excel_files(folder)

    if not files:
        print("No .xlsx or .xls files found in the folder.")
        sys.exit(0)

    print(f"Found {len(files)} Excel file(s):")
    for f in files:
        size_mb = os.path.getsize(f) / 1_048_576
        print(f"  - {os.path.basename(f)}  ({size_mb:.1f} MB)")

    # ---- Collect all columns -------------------------------------------
    print("\nScanning column headers across all files...")
    all_columns = _collect_all_columns(files)

    if not all_columns:
        print("No column headers found in any file.")
        sys.exit(1)

    # ---- Column selection ----------------------------------------------
    included_cols = _prompt_exclude_columns(all_columns)
    if not included_cols:
        print("No columns remaining after exclusion — nothing to write.")
        sys.exit(0)

    print(f"\nRetaining {len(included_cols)} column(s).")

    # col_map: lowercase column name -> index in the output row
    col_map = {col.lower(): i for i, col in enumerate(included_cols)}
    n_out = len(included_cols)

    # ---- Filter setup --------------------------------------------------
    filters = _prompt_filters(included_cols)

    # ---- Prepare output workbook ---------------------------------------
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = os.path.join(folder, f"consolidated_{timestamp}.xlsx")

    out_wb = openpyxl.Workbook(write_only=True)
    out_ws = out_wb.create_sheet()
    out_ws.append(included_cols)   # header row

    # ---- Process files -------------------------------------------------
    print(f"\nMerging into: {output_path}\n")

    total_ok = 0
    total_skipped = 0
    total_rows = 0

    for filepath in files:
        fname = os.path.basename(filepath)
        ext = os.path.splitext(filepath)[1].lower()
        file_rows = 0

        print(f"  [ ] {fname}", end="", flush=True)

        try:
            gen = (
                _stream_xlsx(filepath, col_map, n_out, filters)
                if ext == ".xlsx"
                else _stream_xls(filepath, col_map, n_out, filters)
            )

            for out_row in gen:
                out_ws.append(out_row)
                file_rows += 1
                if file_rows % 10_000 == 0:
                    print(
                        f"\r  [~] {fname}  ({file_rows:,} rows so far...)",
                        end="",
                        flush=True,
                    )

            print(f"\r  [✓] {fname}  — {file_rows:,} rows")
            total_ok += 1
            total_rows += file_rows

        except Exception as exc:
            print(f"\r  [!] {fname}  — skipped ({exc})")
            total_skipped += 1

    # ---- Save ----------------------------------------------------------
    print("\nSaving output file...", end="", flush=True)
    out_wb.save(output_path)
    print(" done.")

    # ---- Summary -------------------------------------------------------
    width = 56
    print()
    print("=" * width)
    print("  SUMMARY")
    print("=" * width)
    print(f"  Files processed  : {total_ok}")
    if total_skipped:
        print(f"  Files skipped    : {total_skipped}")
    print(f"  Total rows merged: {total_rows:,}")
    print(f"  Output file      : {output_path}")
    print("=" * width)


if __name__ == "__main__":
    main()
